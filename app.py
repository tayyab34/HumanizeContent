
import os, re, io
from pathlib import Path
import streamlit as st

try:
    import fitz
except:
    fitz = None

try:
    from docx import Document
except:
    Document = None

try:
    from pptx import Presentation
except:
    Presentation = None

try:
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity
except:
    TfidfVectorizer = None
    cosine_similarity = None

st.set_page_config(page_title="Document Integrity RAG", layout="wide")
st.title("Document Integrity RAG (Streamlit)")

def read_pdf(data):
    doc = fitz.open(stream=data, filetype="pdf")
    return "\\n".join(page.get_text("text") for page in doc)

def read_docx(data):
    doc = Document(io.BytesIO(data))
    return "\\n".join(p.text for p in doc.paragraphs)

def read_pptx(data):
    prs = Presentation(io.BytesIO(data))
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                out.append(shape.text)
    return "\\n".join(out)

def extract(uploaded):
    data = uploaded.read()
    ext = Path(uploaded.name).suffix.lower()

    if ext == ".pdf":
        return read_pdf(data)
    if ext == ".docx":
        return read_docx(data)
    if ext == ".pptx":
        return read_pptx(data)

    return data.decode("utf-8", errors="ignore")

def similarity(main_text, refs):
    if not refs:
        return []

    if TfidfVectorizer and cosine_similarity:
        docs = [main_text] + [x[1] for x in refs]
        vec = TfidfVectorizer(stop_words="english")
        mat = vec.fit_transform(docs)
        sims = cosine_similarity(mat[0:1], mat[1:]).ravel() * 100
        return [(refs[i][0], round(float(sims[i]),2)) for i in range(len(refs))]

    return []

main_doc = st.file_uploader("Upload Main Document", type=["pdf","docx","pptx","txt","md","csv"])
ref_docs = st.file_uploader("Upload Reference Documents", accept_multiple_files=True,
                            type=["pdf","docx","pptx","txt","md","csv"])

if st.button("Analyze") and main_doc:
    main_text = extract(main_doc)

    refs = []
    for f in ref_docs or []:
        refs.append((f.name, extract(f)))

    scores = similarity(main_text, refs)

    st.subheader("Analysis Report")
    st.write(f"Characters: {len(main_text):,}")
    st.write(f"Words: {len(main_text.split()):,}")

    if scores:
        st.subheader("Similarity")
        for name, score in sorted(scores, key=lambda x: x[1], reverse=True):
            st.write(f"{name}: {score:.2f}%")

    st.subheader("Extracted Text")
    st.text_area("", main_text[:12000], height=300)
